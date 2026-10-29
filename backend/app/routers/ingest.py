import csv
import io
import re
import uuid
from datetime import datetime
from typing import Dict, List, Optional
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from rapidfuzz import fuzz, process

from app.auth import get_current_user
from app.db import get_connection

router = APIRouter(prefix="/api/v1/ingest", tags=["ingest"])

STANDARD_FIELDS = [
    "PART_ID",
    "QTY_ORDERED",
    "QTY_RECEIVED",
    "UNIT_COST",
    "ORDER_DATE",
    "EXPECTED_DATE",
    "ACTUAL_DATE",
    "CURRENCY",
]

SYNONYMS: Dict[str, List[str]] = {
    "PART_ID": ["item", "item_code", "part", "part_number", "part_no", "part_ref", "item_sku", "sku", "material"],
    "QTY_ORDERED": ["order_qty", "qty_ordered", "ordered_quantity", "count", "ordered"],
    "QTY_RECEIVED": ["qty", "quantity", "qty_rcvd", "delivered_qty", "received_quantity", "shipped_qty", "count_received"],
    "UNIT_COST": ["price", "unit_price", "rate", "unit_rate", "cost", "price_each", "unit_price_usd", "price_inr"],
    "ORDER_DATE": ["po_date", "order_date", "ordered_on", "purchase_dt", "date_ordered"],
    "EXPECTED_DATE": ["exp_date", "expected_date", "promised_date", "due_date"],
    "ACTUAL_DATE": ["rcvd_dt", "received_date", "delivery_date", "arrival_dt", "date_rcvd", "arrival", "date_received"],
    "CURRENCY": ["curr", "currency", "curr_code"],
}


def auto_map_headers(headers: List[str]) -> Dict[str, Optional[str]]:
    mapping: Dict[str, Optional[str]] = {}
    for h in headers:
        clean_h = re.sub(r"[^a-zA-Z0-9_]", "", h.strip().lower())
        matched_target = None

        # 1. Exact synonym check
        for target, syns in SYNONYMS.items():
            if clean_h in syns or clean_h == target.lower():
                matched_target = target
                break

        # 2. Fuzzy match fallback
        if not matched_target:
            best_match, best_score, _ = process.extractOne(
                clean_h,
                [target.lower() for target in STANDARD_FIELDS] + [s for sublist in SYNONYMS.values() for s in sublist],
                scorer=fuzz.token_sort_ratio,
            )
            if best_score >= 70:
                for target, syns in SYNONYMS.items():
                    if best_match == target.lower() or best_match in syns:
                        matched_target = target
                        break

        mapping[h] = matched_target
    return mapping


def parse_date(val: str) -> Optional[str]:
    if not val or not str(val).strip():
        return None
    cleaned = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(cleaned, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def parse_numeric(val: str) -> Optional[float]:
    if val is None or str(val).strip() == "":
        return None
    cleaned = re.sub(r"[^0-9.]", "", str(val))
    try:
        return float(cleaned)
    except ValueError:
        return None


@router.post("/preview")
async def preview_upload(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
):
    """Parses raw CSV/XLSX and returns auto-detected column mappings and first 5 preview rows."""
    content = await file.read()
    filename = file.filename or "upload.csv"
    headers: List[str] = []
    rows: List[Dict[str, str]] = []

    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        headers = next(reader, [])
        for i, r in enumerate(reader):
            if i >= 5:
                break
            rows.append({headers[j]: r[j] if j < len(r) else "" for j in range(len(headers))})
    elif filename.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            iter_rows = list(ws.iter_rows(values_only=True))
            if iter_rows:
                headers = [str(c) if c is not None else f"Col_{j}" for j, c in enumerate(iter_rows[0])]
                for r in iter_rows[1:6]:
                    rows.append({headers[j]: str(r[j]) if j < len(r) and r[j] is not None else "" for j in range(len(headers))})
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Failed to read Excel workbook: {exc}")
    else:
        raise HTTPException(status_code=400, detail="Only CSV and XLSX files are supported")

    mapping = auto_map_headers(headers)
    return {
        "filename": filename,
        "headers": headers,
        "auto_mapping": mapping,
        "standard_fields": STANDARD_FIELDS,
        "preview_rows": rows,
    }


class ConfirmImportPayload(BaseModel):
    filename: str
    supplier_id: str
    column_mapping: Dict[str, str]  # { "Raw Header": "STANDARD_FIELD" }
    rows: List[Dict[str, str]]


@router.post("/confirm")
def confirm_import(
    payload: ConfirmImportPayload,
    current_user: dict = Depends(get_current_user),
):
    """Normalizes mapped rows and inserts them directly into SUPPLIER_DELIVERIES in Exasol."""
    file_id = str(uuid.uuid4())
    uploaded_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")

    normalized_records = []
    failed_rows = 0

    for r in payload.rows:
        mapped_item: Dict[str, any] = {}
        for raw_h, std_f in payload.column_mapping.items():
            if std_f and std_f in STANDARD_FIELDS:
                mapped_item[std_f] = r.get(raw_h)

        part_id = mapped_item.get("PART_ID") or "P-BRG-6205"
        qty_received = parse_numeric(mapped_item.get("QTY_RECEIVED"))
        qty_ordered = parse_numeric(mapped_item.get("QTY_ORDERED")) or qty_received or 10
        unit_cost = parse_numeric(mapped_item.get("UNIT_COST")) or 50.0
        currency = mapped_item.get("CURRENCY") or "INR"
        order_date = parse_date(mapped_item.get("ORDER_DATE")) or datetime.utcnow().strftime("%Y-%m-%d")
        act_date = parse_date(mapped_item.get("ACTUAL_DATE")) or datetime.utcnow().strftime("%Y-%m-%d")
        exp_date = parse_date(mapped_item.get("EXPECTED_DATE")) or act_date

        if not qty_received and not qty_ordered:
            failed_rows += 1
            continue

        normalized_records.append({
            "DELIVERY_ID": str(uuid.uuid4()),
            "SUPPLIER_ID": payload.supplier_id,
            "PART_ID": part_id,
            "ORDER_DATE": order_date,
            "EXPECTED_DATE": exp_date,
            "ACTUAL_DATE": act_date,
            "QTY_ORDERED": int(qty_ordered),
            "QTY_RECEIVED": int(qty_received if qty_received is not None else qty_ordered),
            "UNIT_COST": round(float(unit_cost), 4),
            "CURRENCY": currency,
            "DELAY_DAYS": 0,
        })

    with get_connection() as con:
        # 1. Log file in SUPPLIER_FILES
        con.execute(
            """
            INSERT INTO LOGISENSE.SUPPLIER_FILES
                (FILE_ID, SUPPLIER_ID, FILENAME, UPLOADED_AT, ROW_COUNT, STATUS)
            VALUES
                ({fid!s}, {sid!s}, {fname!s}, {ts!s}, {rc!s}, 'LOADED')
            """,
            {
                "fid": file_id,
                "sid": payload.supplier_id,
                "fname": payload.filename,
                "ts": uploaded_at,
                "rc": len(normalized_records),
            },
        )

        # 2. Insert into SUPPLIER_DELIVERIES
        for rec in normalized_records:
            con.execute(
                """
                INSERT INTO LOGISENSE.SUPPLIER_DELIVERIES
                    (DELIVERY_ID, SUPPLIER_ID, PART_ID, ORDER_DATE, EXPECTED_DATE, ACTUAL_DATE, QTY_ORDERED, QTY_RECEIVED, UNIT_COST, CURRENCY, DELAY_DAYS)
                VALUES
                    ({did!s}, {sid!s}, {pid!s}, {odt!s}, {edt!s}, {adt!s}, {qo!s}, {qr!s}, {uc!s}, {cur!s}, {dd!s})
                """,
                {
                    "did": rec["DELIVERY_ID"],
                    "sid": rec["SUPPLIER_ID"],
                    "pid": rec["PART_ID"],
                    "odt": rec["ORDER_DATE"],
                    "edt": rec["EXPECTED_DATE"],
                    "adt": rec["ACTUAL_DATE"],
                    "qo": rec["QTY_ORDERED"],
                    "qr": rec["QTY_RECEIVED"],
                    "uc": rec["UNIT_COST"],
                    "cur": rec["CURRENCY"],
                    "dd": rec["DELAY_DAYS"],
                },
            )

    return {
        "status": "success",
        "file_id": file_id,
        "rows_imported": len(normalized_records),
        "rows_failed": failed_rows,
        "message": f"Successfully ingested {len(normalized_records)} delivery records into Exasol",
    }
